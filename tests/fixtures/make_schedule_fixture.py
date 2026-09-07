"""Synthetic workbook for parser regression tests; contains no production data."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Color

wb = Workbook()
sh = wb.active
sh.title = "本校教務部用"
sh.merge_cells("C3:E3")
sh["C3"] = "小学部 4:55～6:15"
sh.merge_cells("C4:E4")
sh["C4"] = "中学部 6:35～8:05"
for address, value in {"C5": "①", "D5": "②", "E5": "③", "A7": 7, "B7": "月", "C7": "１Ｓ英", "D7": "５Ａ算\n対面", "E7": "英検対策①", "I56": "テスト甲", "J56": "テスト乙", "K56": "テスト丙"}.items():
    sh[address] = value
for cell, legend, color in [("C7", "I56", Color(indexed=42)), ("D7", "J56", Color(rgb="FFFFFF00")), ("E7", "K56", Color(theme=7, tint=-0.5))]:
    sh[cell].fill = PatternFill(patternType="solid", fgColor=color)
    sh[legend].fill = PatternFill(patternType="solid", fgColor=color)
wb.save(Path(__file__).with_name("schedule-synthetic.xlsx"))
