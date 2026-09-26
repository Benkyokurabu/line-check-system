"""Local-only bridge between the authenticated staff page and the NAS sources."""
from __future__ import annotations

import json
import hashlib
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import unicodedata
import uuid
from concurrent.futures import ThreadPoolExecutor
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import fitz
from PIL import Image
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

PORT = 38473
ORIGIN = os.environ.get('BENTAN_ORIGIN', 'https://line-check-system.vercel.app')
ROOT = Path(sys.executable).resolve().parent if getattr(sys, 'frozen', False) else Path(__file__).resolve().parent
CONFIG = ROOT / 'sources.txt'
GUIDE_CONFIG = ROOT / 'guide-path.txt'
SUFFIXES = {'.pdf', '.jpg', '.jpeg', '.png'}
SESSIONS: dict[str, tuple[float, tempfile.TemporaryDirectory, dict, str]] = {}
LOCK = threading.Lock()
OCR_CACHE: dict[str, tuple[int, float, str]] = {}
INDEX_FILE = ROOT / 'hokushin-index.json'
INDEX_STATUS = {'running': False, 'completed': 0, 'total': 0, 'year': ''}


def save_bundle_to_onedrive(source: Path, student_number: str) -> Path:
    onedrive = os.environ.get('OneDrive') or os.environ.get('OneDriveConsumer') or os.environ.get('OneDriveCommercial')
    if not onedrive or not Path(onedrive).is_dir():
        raise RuntimeError('このPCのOneDriveフォルダが見つかりません。OneDriveの接続を確認してください。')
    destination = Path(onedrive) / '面談準備' / '保存済み資料'
    destination.mkdir(parents=True, exist_ok=True)
    filename = f'{time.strftime("%Y%m%d_%H%M%S")}_{student_number}_面談資料_{uuid.uuid4().hex[:8]}.pdf'
    saved = destination / filename
    partial = destination / f'.{filename}.partial'
    try:
        with source.open('rb') as incoming, partial.open('xb') as outgoing:
            shutil.copyfileobj(incoming, outgoing)
        os.replace(partial, saved)
    finally:
        partial.unlink(missing_ok=True)
    return saved


def sync_bundle_to_cloud(saved: Path) -> bool:
    rclone = shutil.which('rclone')
    if not rclone:
        return False  # The OneDrive desktop app can sync the local folder instead.
    try:
        remotes = subprocess.run([rclone, 'listremotes'], capture_output=True, text=True, timeout=10,
                                 creationflags=subprocess.CREATE_NO_WINDOW)
        if remotes.returncode or 'onedrive:' not in remotes.stdout.splitlines():
            return False
        upload = subprocess.run([rclone, 'copyto', str(saved), f'onedrive:面談準備/保存済み資料/{saved.name}'],
                                capture_output=True, text=True, timeout=300,
                                creationflags=subprocess.CREATE_NO_WINDOW)
        if upload.returncode:
            raise RuntimeError('OneDriveのクラウドへ送れませんでした。このPCの保存済みPDFは残っています。もう一度「OneDriveに保存」を押してください。')
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError('OneDriveのクラウドへの送信が時間切れになりました。もう一度「OneDriveに保存」を押してください。') from exc
    return True


def norm(value: str) -> str:
    return re.sub(r'[\s\u3000]', '', unicodedata.normalize('NFKC', str(value))).lower()


def canonical_school_name(value: str) -> str:
    name = unicodedata.normalize('NFKC', str(value)).strip()
    return '叡明' if norm(name).replace('高等学校', '').replace('高校', '') == 'えいめい' else name


def school_key(value: str) -> str:
    value = canonical_school_name(value)
    value = re.sub(r'^[ぁ-ん](?=[\s　]|[一-龯])', '', str(value))
    return norm(value).replace('高等学校', '').replace('高校', '')


def school_name_matches(filename: str, school: str) -> bool:
    stem = school_key(filename)
    key = school_key(school)
    if len(key) < 2:
        return False
    for match in re.finditer(re.escape(key), stem):
        rest = stem[match.end():]
        if not rest or rest[0] in '（([【・_-':
            return True
        if rest.startswith(('推薦', '選抜', '入試', '募集', '基準', '案内', '偏差値')):
            return True
    return False


def roots() -> list[Path]:
    lines = [line.strip() for line in CONFIG.read_text(encoding='utf-8-sig').splitlines() if line.strip()]
    if len(lines) != 9 or any(not line.startswith('\\\\TS3210\\benko\\') for line in lines):
        raise ValueError('資料の場所はNASの9行を指定してください。')
    return [Path(line) for line in lines]


def files(root: Path):
    if not root.is_dir():
        return []
    found = []
    for directory, dirs, names in os.walk(root):
        dirs[:] = [d for d in dirs if not d.startswith('●作業用') and not d.startswith('過去資料')]
        for name in names:
            path = Path(directory) / name
            if path.suffix.lower() in SUFFIXES:
                found.append(path)
    return found


def display_label(path: Path, root: Path) -> str:
    return str(path.relative_to(root)).replace('\\', ' / ')


def source_pdf(path: Path, output: Path) -> Path:
    if path.suffix.lower() == '.pdf':
        return path
    image = Image.open(path)
    image.load()
    image = image.convert('RGB')
    image.save(output, 'PDF', resolution=180.0)
    return output


def selected_schools(all_roots: list[Path], names: list[str]):
    selected = []
    missing = []
    indexed = {source_id: files(all_roots[source_id]) for source_id in (3, 4, 5, 6)}
    for school in names:
        key = school_key(school)
        if len(key) < 2:
            missing.append(f'{school}：学校名を確認してください')
            continue
        found_for_school = []
        for source_id in (3, 4, 5, 6):
            candidates = indexed[source_id]
            matches = []
            for path in candidates:
                if not school_name_matches(path.stem, school):
                    continue
                matches.append(path)
            if source_id == 5 and matches:
                years = [(int(match[1]), path) for path in matches if (match := re.search(r'(20\d{2})年受験用', str(path)))]
                matches = [path for year, path in years if year == max(y for y, _ in years)] if years else []
            found_for_school.extend((source_id, path) for path in matches)
        if not found_for_school:
            missing.append(f'{school}：学校資料が見つかりません')
        selected.extend(found_for_school)
    return list(dict.fromkeys(selected)), missing


def preview_schools(all_roots: list[Path], names: list[str]) -> list[dict]:
    if not names:
        return []
    indexed = {source_id: files(all_roots[source_id]) for source_id in (3, 4, 5, 6)}
    result = []
    labels = {3: '高校案内', 4: '選抜基準', 5: '私立推薦基準', 6: '北辰偏差値資料'}
    for rank, raw_name in enumerate(names, 1):
        name = canonical_school_name(raw_name)
        found = []
        for source_id, candidates in indexed.items():
            matches = [path for path in candidates if school_name_matches(path.stem, name)]
            if source_id == 5 and matches:
                years = [(int(match[1]), path) for path in matches if (match := re.search(r'(20\d{2})年受験用', str(path)))]
                matches = [path for year, path in years if year == max(y for y, _ in years)] if years else []
            for path in matches:
                year_match = re.search(r'(20\d{2})年(?:度|受験用)?', str(path)) or re.search(r'(20\d{2})', str(path))
                found.append({'kind': labels[source_id], 'year': f'{year_match[1]}年度' if year_match else '年度不明',
                              'filename': path.name})
        result.append({'rank': rank, 'surveyName': raw_name, 'name': name, 'found': bool(found), 'files': found})
    return result


def latest_year_root(configured: Path) -> Path:
    siblings = []
    for path in configured.parent.iterdir():
        match = re.fullmatch(r'(20\d{2})年度', path.name)
        if path.is_dir() and match:
            siblings.append((int(match[1]), path))
    for _, path in sorted(siblings, reverse=True):
        if next(path.rglob('*.pdf'), None) is not None:
            return path
    return configured


def tesseract_path() -> Path | None:
    for candidate in (ROOT / 'ocr' / 'tesseract.exe', Path(r'C:\Program Files\Tesseract-OCR\tesseract.exe')):
        if candidate.is_file():
            return candidate
    command = shutil.which('tesseract')
    return Path(command) if command else None


def ocr_pdf(path: Path) -> str:
    stat = path.stat()
    key = str(path)
    cached = OCR_CACHE.get(key)
    if cached and cached[:2] == (stat.st_size, stat.st_mtime):
        return cached[2]
    tesseract = tesseract_path()
    if not tesseract:
        return ''
    with fitz.open(str(path)) as document:
        if not document.page_count:
            return ''
        image = document[0].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).tobytes('png')
    result = subprocess.run([str(tesseract), 'stdin', 'stdout', '-l', 'jpn+eng', '--psm', '11'],
                            input=image, capture_output=True, timeout=45,
                            creationflags=subprocess.CREATE_NO_WINDOW)
    content = result.stdout.decode('utf-8', errors='replace') if result.returncode == 0 else ''
    OCR_CACHE[key] = (stat.st_size, stat.st_mtime, content)
    return content


def hokushin_rank(path: Path, year: int) -> tuple[int, int, int, int]:
    folder = path.parent.name
    date = re.search(r'(\d{2})月(\d{2})日', folder)
    round_match = re.search(r'第0?(\d+)回', folder)
    return (year, int(date[1]) if date else 0, int(date[2]) if date else 0,
            int(round_match[1]) if round_match else 0)


def read_hokushin_index() -> dict:
    try:
        return json.loads(INDEX_FILE.read_text(encoding='utf-8'))
    except (OSError, ValueError):
        return {}


def write_hokushin_index(year: str, records: list[dict], complete: bool) -> None:
    INDEX_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary = INDEX_FILE.with_suffix('.partial')
    temporary.write_text(json.dumps({'year': year, 'complete': complete, 'records': records}, ensure_ascii=False),
                         encoding='utf-8')
    os.replace(temporary, INDEX_FILE)


def build_hokushin_index() -> None:
    if not tesseract_path():
        raise RuntimeError('日本語OCR（Tesseract）がこのPCにありません')
    year_root = latest_year_root(roots()[1])
    year_match = re.search(r'20\d{2}', year_root.name)
    year = int(year_match[0]) if year_match else 0
    paths = [path for path in files(year_root) if path.suffix.lower() == '.pdf' and
             '個人成績表' in str(path.parent)]
    paths.sort(key=lambda path: (hokushin_rank(path, year), '南教室' in path.parent.name), reverse=True)
    existing = read_hokushin_index()
    previous = {record['source']: record for record in existing.get('records', [])} if existing.get('year') == year_root.name else {}
    INDEX_STATUS.update({'running': True, 'completed': 0, 'total': len(paths), 'year': year_root.name})
    records = []
    cache_dir = ROOT / 'hokushin-cache' / year_root.name
    cache_dir.mkdir(parents=True, exist_ok=True)

    def index_one(path: Path) -> dict:
        stat = path.stat()
        old = previous.get(str(path))
        if old and old.get('size') == stat.st_size and old.get('mtimeNs') == stat.st_mtime_ns and Path(old['local']).is_file():
            filename_key = norm(path.stem)
            return {**old, 'text': old['text'] if filename_key in old['text'] else old['text'] + filename_key}
        local = cache_dir / (hashlib.sha256(str(path).encode('utf-8')).hexdigest()[:24] + '.pdf')
        partial = local.with_suffix('.partial')
        try:
            shutil.copy2(path, partial)
            os.replace(partial, local)
        finally:
            partial.unlink(missing_ok=True)
        content = ''
        for text_path in (path.with_name(path.stem + '_ocr.txt'), path.with_suffix('.txt')):
            if text_path.is_file():
                content = text_path.read_text(encoding='utf-8-sig', errors='ignore')
                break
        if not content:
            with fitz.open(str(local)) as document:
                content = document[0].get_text() if document.page_count else ''
            if len(content.strip()) < 30:
                content = ocr_pdf(local)
        return {'source': str(path), 'local': str(local), 'size': stat.st_size, 'mtimeNs': stat.st_mtime_ns,
                'rank': hokushin_rank(path, year), 'grade': '中3' if '中3' in norm(path.parent.name) else '中2',
                'campus': '南教室' if '南教室' in path.parent.name else '本校', 'text': norm(content + ' ' + path.stem)}

    def safe_index_one(path: Path) -> dict:
        try:
            return index_one(path)
        except Exception as exc:
            return {'source': str(path), 'grade': '中3' if '中3' in norm(path.parent.name) else '中2',
                    'error': str(exc)[:160]}

    try:
        with ThreadPoolExecutor(max_workers=3) as pool:
            for index, record in enumerate(pool.map(safe_index_one, paths), 1):
                records.append(record)
                INDEX_STATUS['completed'] = index
                if index % 10 == 0:
                    write_hokushin_index(year_root.name, records, False)
        write_hokushin_index(year_root.name, records, True)
    finally:
        INDEX_STATUS['running'] = False


def indexed_hokushin(year_root: Path, grade: str, name: str, number: str, campus: str) -> tuple[Path | None, str | None]:
    index = read_hokushin_index()
    if INDEX_STATUS.get('error'):
        return None, f'北辰：索引準備に失敗しました：{INDEX_STATUS["error"]}'
    if INDEX_STATUS['running'] or index.get('year') != year_root.name or not index.get('complete'):
        done, total = INDEX_STATUS['completed'], INDEX_STATUS['total']
        return None, f'北辰：事前索引を作成中です（{done}/{total}件）。完了後に資料を再確認してください'
    target = norm(name)
    matches = [record for record in index['records'] if record.get('grade') == grade and 'text' in record and
               ((target and target in record['text']) or (number and number in re.sub(r'\D', '', record['text'])))]
    matches.sort(key=lambda record: (record['rank'], record['campus'] == campus), reverse=True)
    if not matches:
        if any(record.get('grade') == grade and record.get('error') for record in index['records']):
            return None, '北辰：一部の成績票を読めないため、生徒の票を特定できません'
        return None, '北辰：個人成績票を特定できません'
    best = matches[0]
    ambiguous = [record for record in matches if record['rank'] == best['rank'] and record['campus'] == best['campus']]
    if len(ambiguous) != 1:
        return None, '北辰：同じ最新回に複数の候補があるため選べません'
    return Path(best['local']), None


def hokushin_source(path: Path) -> Path:
    for record in read_hokushin_index().get('records', []):
        if record.get('local') == str(path):
            return Path(record['source'])
    return path


def hokushin(all_roots: list[Path], grade: str, name: str, number: str = '', campus: str = ''):
    target = norm(name)
    year_root = latest_year_root(all_roots[1])
    if getattr(sys, 'frozen', False):
        return indexed_hokushin(year_root, grade, name, number, campus)
    year = int(re.search(r'20\d{2}', year_root.name)[0]) if re.search(r'20\d{2}', year_root.name) else 0
    paths = [path for path in files(year_root) if path.suffix.lower() == '.pdf' and
             '個人成績表' in str(path.parent) and norm(grade) in norm(path.parent.name)]
    paths.sort(key=lambda path: hokushin_rank(path, year), reverse=True)

    def is_match(content: str) -> bool:
        normalized = norm(content)
        return bool((target and target in normalized) or (number and number in re.sub(r'\D', '', content)))

    # Search one round at a time so a newer PDF without a sidecar wins over
    # an older PDF with one.
    for rank in sorted({hokushin_rank(path, year) for path in paths}, reverse=True):
        round_paths = [path for path in paths if hokushin_rank(path, year) == rank]
        groups = [round_paths]
        if campus in ('本校', '南教室'):
            preferred = [path for path in round_paths if campus in path.parent.name]
            others = [path for path in round_paths if campus not in path.parent.name]
            groups = [group for group in (preferred, others) if group]
        for group in groups:
            sidecar_matches = []
            for path in group:
                for text_path in (path.with_name(path.stem + '_ocr.txt'), path.with_suffix('.txt')):
                    if not text_path.is_file():
                        continue
                    try:
                        if is_match(text_path.read_text(encoding='utf-8-sig', errors='ignore')):
                            sidecar_matches.append(path)
                    except OSError:
                        pass
                    break
            if len(sidecar_matches) == 1:
                return sidecar_matches[0], None
            if len(sidecar_matches) > 1:
                return None, '北辰：同じ最新回に複数の候補があるため選べません'
            def read_page(path: Path) -> tuple[Path, str]:
                try:
                    with fitz.open(str(path)) as document:
                        text = document[0].get_text() if document.page_count else ''
                    if not is_match(text):
                        text = ocr_pdf(path)
                    return path, text
                except Exception:
                    return path, ''
            with ThreadPoolExecutor(max_workers=4) as pool:
                matches = [path for path, content in pool.map(read_page, group) if is_match(content)]
            if len(matches) == 1:
                return matches[0], None
            if len(matches) > 1:
                return None, '北辰：同じ最新回に複数の候補があるため選べません'
    return None, '北辰：個人成績票を特定できません'


def preview_hokushin(all_roots: list[Path], grade: str, name: str, number: str, campus: str = '') -> dict:
    if grade not in ('中2', '中3'):
        return {'found': False, 'message': '対象学年ではありません'}
    path, error = hokushin(all_roots, grade, name, number, campus)
    if not path:
        return {'found': False, 'indexing': bool(error and ('索引を作成中' in error or '索引準備に失敗' in error)),
                'message': error}
    source = hokushin_source(path)
    year = re.search(r'20\d{2}', str(source))
    return {'found': True, 'year': f'{year[0]}年度' if year else '年度不明',
            'round': source.parent.name, 'filename': source.name}


def report_periods(configured: Path) -> list[tuple[int, int, str, Path]]:
    """Find the student's current and past report folders beside the configured term."""
    parent = configured.parent.parent
    if not parent.is_dir():
        return []
    periods = []
    for folder in parent.iterdir():
        match = re.search(r'(20\d{2})年度(前期|後期)成績通知', folder.name)
        if folder.is_dir() and match:
            report_folder = folder / configured.name
            if report_folder.is_dir():
                periods.append((int(match[1]), 1 if match[2] == '後期' else 0, match[2], report_folder))
    return sorted(periods, reverse=True)


def local_report(path: Path) -> Path:
    """Keep large class PDFs on this PC so repeated student lookups do not reread NAS pages."""
    stat = path.stat()
    cache = ROOT / 'report-cache'
    cache.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(str(path).encode('utf-8')).hexdigest()[:20]
    local = cache / f'{digest}{path.suffix.lower()}'
    if not local.is_file() or local.stat().st_size != stat.st_size or local.stat().st_mtime != stat.st_mtime:
        partial = cache / f'{digest}.partial'
        try:
            shutil.copy2(path, partial)
            os.replace(partial, local)
        finally:
            partial.unlink(missing_ok=True)
    return local


def report_pages(path: Path, name: str, number: str) -> list[int]:
    local = local_report(path)
    name_key = norm(name)
    if local.suffix.lower() != '.pdf':
        return [0] if (number and number in norm(path.stem)) or (name_key and name_key in norm(path.stem)) else []
    pages = []
    with fitz.open(str(local)) as document:
        scanned_pages = []
        for index, page in enumerate(document):
            content = norm(page.get_text())
            if name_key and name_key in content or number and number in content:
                pages.append(index)
            elif len(content) < 50:
                scanned_pages.append(index)
        if pages:
            return pages
        if document.page_count == 1 and ((number and number in norm(path.stem)) or (name_key and name_key in norm(path.stem))):
            return [0]
        tesseract = tesseract_path()
        if not tesseract:
            return []
        for index in scanned_pages:
            page = document[index]
            image = page.get_pixmap(matrix=fitz.Matrix(1.25, 1.25), alpha=False).tobytes('png')
            result = subprocess.run([str(tesseract), 'stdin', 'stdout', '-l', 'jpn+eng', '--psm', '11'],
                                    input=image, capture_output=True, timeout=45,
                                    creationflags=subprocess.CREATE_NO_WINDOW)
            content = norm(result.stdout.decode('utf-8', errors='replace')) if result.returncode == 0 else ''
            if name_key and name_key in content or number and number in content:
                pages.append(index)
    return pages


def report_candidate_rank(path: Path, grade: str, campus: str) -> tuple[int, int, float]:
    label = norm(path.stem)
    grade_label = norm(grade)
    explicit_grade = bool(re.search(r'(?:中|小)[1-6]', label))
    if explicit_grade and grade_label not in label:
        return (-1, -1, 0)
    campus_label = '南' if '南' in campus else '本' if '本' in campus else ''
    explicit_campus = '南' if '南' in label else '本' if '本校' in label else ''
    campus_score = 1 if campus_label and explicit_campus == campus_label else -1 if campus_label and explicit_campus and explicit_campus != campus_label else 0
    return (1 if explicit_grade else 0, campus_score, path.stat().st_mtime)


def term_report(all_roots: list[Path], grade: str, name: str, number: str, campus: str = '') -> tuple[Path | None, list[int], str | None]:
    if not norm(name) or not re.fullmatch(r'\d{5,12}', number):
        return None, [], '成績通知：生徒情報を確認してください'
    source_id = 7 if grade.startswith('中') else 8 if grade.startswith('小') else None
    if source_id is None:
        return None, [], '成績通知：対象学年ではありません'
    for _, _, _, folder in report_periods(all_roots[source_id]):
        candidates = sorted(files(folder), key=lambda path: report_candidate_rank(path, grade, campus), reverse=True)
        for path in candidates:
            grade_score, campus_score, _ = report_candidate_rank(path, grade, campus)
            if grade_score < 0 or campus_score < 0:
                continue
            pages = report_pages(path, name, number)
            if pages:
                return path, pages, None
    return None, [], '成績通知：本人の個人成績表が見つかりません'


def preview_term_report(all_roots: list[Path], grade: str, name: str, number: str, campus: str = '') -> dict:
    path, pages, error = term_report(all_roots, grade, name, number, campus)
    if not path:
        return {'found': False, 'message': error}
    match = re.search(r'(20\d{2})年度(前期|後期)成績通知', str(path))
    return {'found': True, 'year': match[1] if match else '', 'term': match[2] if match else '',
            'filename': path.name, 'pages': [page + 1 for page in pages]}


def extract_report_pages(path: Path, pages: list[int], output: Path) -> Path:
    if path.suffix.lower() != '.pdf':
        return source_pdf(path, output)
    with fitz.open(str(local_report(path))) as source, fitz.open() as selected:
        for page in pages:
            selected.insert_pdf(source, from_page=page, to_page=page)
        selected.save(str(output))
    return output


def vmogi(all_roots: list[Path], number: str, name: str, grade: str):
    root = all_roots[2]
    books = list(root.glob('*.xlsx'))
    books.sort(key=lambda p: int(re.search(r'(\d+)月号', p.name)[1]) if re.search(r'(\d+)月号', p.name) else 0, reverse=True)
    for book in books:
        wb = load_workbook(book, read_only=True, data_only=True)
        try:
            sheet = wb.active
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value or '').strip() for value in next(rows)]
            index = {header: pos for pos, header in enumerate(headers)}
            if '氏名' not in index:
                continue
            matches = []
            for row in rows:
                if norm(row[index['氏名']]) != norm(name):
                    continue
                if '学年' in index and str(row[index['学年']] or '').strip() not in (grade[-1], f'中{grade[-1]}'):
                    continue
                if '登録番号' in index and number and str(row[index['登録番号']] or '').strip() == number:
                    return book, headers, row, None
                matches.append(row)
            if len(matches) == 1:
                return book, headers, matches[0], None
            if len(matches) > 1:
                return None, None, None, 'Vもぎ：同姓同名の候補が複数あります'
        finally:
            wb.close()
    return None, None, None, 'Vもぎ：成績を特定できません'


def vmogi_pdf(book: Path, headers: list[str], row: tuple, output: Path):
    pdfmetrics.registerFont(UnicodeCIDFont('HeiseiKakuGo-W5'))
    data = [(headers[i], str(value)) for i, value in enumerate(row) if value is not None and headers[i]]
    c = canvas.Canvas(str(output), pagesize=A4)
    width, height = A4
    c.setFont('HeiseiKakuGo-W5', 16)
    c.drawString(38, height - 42, 'Vもぎ 個人成績')
    c.setFont('HeiseiKakuGo-W5', 9)
    c.drawString(38, height - 60, book.name)
    y = height - 82
    for label, value in data:
        if y < 42:
            c.showPage()
            c.setFont('HeiseiKakuGo-W5', 9)
            y = height - 42
        c.drawString(38, y, label[:22])
        c.drawString(190, y, value[:76])
        c.setStrokeColor(colors.lightgrey)
        c.line(38, y - 4, width - 38, y - 4)
        y -= 18
    c.save()


def guide_pdf(number: str, name: str, output: Path):
    script = ROOT / 'export-guide.ps1'
    master = GUIDE_CONFIG.read_text(encoding='utf-8-sig').strip()
    if not master.startswith('\\\\TS3210\\benko\\') or not master.lower().endswith('.xlsm'):
        raise ValueError('指導簿の場所を確認してください。')
    result = subprocess.run(['powershell.exe', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(script),
                             '-StudentNumber', number, '-StudentName', name, '-MasterPath', master, '-OutputPath', str(output)],
                            capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=120,
                            creationflags=subprocess.CREATE_NO_WINDOW)
    if result.returncode or not output.is_file():
        raise RuntimeError('指導簿の生徒照合またはPDF化に失敗しました')


def make_bundle(payload: dict):
    number = str(payload.get('number', ''))
    name = str(payload.get('name', '')).strip()
    grade = unicodedata.normalize('NFKC', str(payload.get('grade', '')))
    schools = payload.get('schools', [])
    if not re.fullmatch(r'\d{5,12}', number) or not name or len(name) > 80 or not re.fullmatch(r'(小[4-6]|中[1-3])', grade):
        raise ValueError('生徒情報を確認してください。')
    if not isinstance(schools, list) or len(schools) > 6 or any(not isinstance(s, str) or len(s) > 80 for s in schools):
        raise ValueError('志望校を確認してください。')
    all_roots = roots()
    folder = tempfile.TemporaryDirectory(prefix='bentan-materials-')
    base = Path(folder.name)
    items = []
    missing = []
    pdfs = []

    def add(label, source, sensitive=False):
        index = len(items)
        pdf = source_pdf(source, base / f'converted-{index}.pdf')
        PdfReader(str(pdf))  # Fail a damaged source before returning a successful bundle.
        pdfs.append(pdf)
        items.append({'label': label, 'source': str(source), 'staffOnly': sensitive})

    try:
        guide = base / 'guide.pdf'
        try:
            guide_pdf(number, name, guide)
            add('指導簿', guide)
        except Exception as exc:
            missing.append(str(exc))
        if grade == '中3':
            common = [p for p in files(all_roots[0]) if p.parent == all_roots[0] and p.suffix.lower() == '.pdf']
            for path in common:
                if any(term in norm(path.stem) for term in ('偏差値段階表', '偏差値基準')):
                    continue
                add('中3共通資料：' + path.name, path)
        school_files, school_missing = selected_schools(all_roots, schools)
        missing.extend(school_missing)
        for source_id, path in school_files:
            sensitive = source_id == 5
            label = f'{["", "", "", "高校案内", "選抜基準", "私立推薦基準", "北辰偏差値資料"][source_id]}：{path.name}'
            try:
                add(label, path, sensitive)
            except Exception:
                missing.append(label + '：PDF化できません')
        if grade in ('中2', '中3'):
            path, error = hokushin(all_roots, grade, name, number, str(payload.get('campus', '')))
            if path:
                add('北辰：' + hokushin_source(path).parent.name, path)
            elif error:
                if '索引を作成中' in error or '索引準備に失敗' in error:
                    raise RuntimeError(error)
                missing.append(error)
            book, headers, row, error = vmogi(all_roots, number, name, grade)
            if book:
                target = base / 'vmogi.pdf'
                vmogi_pdf(book, headers, row, target)
                add('Vもぎ：' + book.name, target)
            elif error:
                missing.append(error)
        report, report_pages_found, report_error = term_report(all_roots, grade, name, number, str(payload.get('campus', '')))
        if report:
            try:
                selected_report = extract_report_pages(report, report_pages_found, base / 'term-report.pdf')
                add('成績通知：' + report.parent.parent.name + '／' + report.name + '（本人のページ）', selected_report)
            except Exception:
                missing.append('成績通知：個人成績表をPDFにできませんでした')
        elif report_error:
            missing.append(report_error)
        if not pdfs:
            raise RuntimeError('印刷できる資料が見つかりません')
        writer = PdfWriter()
        for pdf in pdfs:
            writer.append(str(pdf))
        combined = base / 'staff-bundle.pdf'
        with combined.open('wb') as stream:
            writer.write(stream)
        return folder, {'items': items, 'missing': missing, 'hasGuide': guide.is_file(), 'pages': len(writer.pages)}
    except Exception:
        folder.cleanup()
        raise


class Handler(BaseHTTPRequestHandler):
    def cors(self):
        self.send_header('Access-Control-Allow-Origin', ORIGIN)
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.send_header('Access-Control-Allow-Private-Network', 'true')
        self.send_header('Vary', 'Origin')
        self.send_header('Cache-Control', 'no-store')

    def allowed(self):
        return self.headers.get('Origin') == ORIGIN and self.headers.get('Host') in (f'127.0.0.1:{PORT}', f'localhost:{PORT}')

    def reply(self, status, body):
        raw = json.dumps(body, ensure_ascii=False).encode('utf-8')
        self.send_response(status)
        self.cors()
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Content-Length', str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)

    def do_OPTIONS(self):
        if not self.allowed():
            self.send_error(403)
            return
        self.send_response(204)
        self.cors()
        self.end_headers()

    def do_GET(self):
        if self.path == '/health':
            self.reply(200, {'ready': True, 'version': 2, 'hokushinIndex': INDEX_STATUS})
            return
        path, _, query = self.path.partition('?')
        match = re.fullmatch(r'/pdf/([a-f0-9]{32})/(staff-bundle|guide)\.pdf', path)
        if not match:
            self.send_error(404)
            return
        with LOCK:
            session = SESSIONS.get(match[1])
        if not session or session[0] < time.time():
            self.send_error(404)
            return
        path = Path(session[1].name) / (match[2] + '.pdf')
        if not path.is_file():
            self.send_error(404)
            return
        data = path.read_bytes()
        self.send_response(200)
        self.cors()
        self.send_header('Content-Type', 'application/pdf')
        disposition = 'attachment' if query == 'download=1' else 'inline'
        self.send_header('Content-Disposition', f'{disposition}; filename="{match[2]}.pdf"')
        self.send_header('Content-Length', str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_POST(self):
        if self.path == '/preview':
            if not self.allowed() or self.headers.get('Content-Type', '').split(';')[0] != 'application/json':
                self.send_error(403)
                return
            try:
                length = int(self.headers.get('Content-Length', '0'))
                if length < 1 or length > 4096:
                    raise ValueError('入力が大きすぎます')
                payload = json.loads(self.rfile.read(length))
                names = payload.get('schools')
                if not isinstance(names, list) or len(names) > 6 or any(not isinstance(s, str) or len(s) > 80 for s in names):
                    raise ValueError('志望校を確認してください')
                all_roots = roots()
                grade, name, number = (str(payload.get(key, '')) for key in ('grade', 'name', 'number'))
                campus = str(payload.get('campus', ''))
                INDEX_STATUS['previewStage'] = 'schools'
                schools = preview_schools(all_roots, names)
                INDEX_STATUS['previewStage'] = 'hokushin'
                north = preview_hokushin(all_roots, grade, name, number, campus)
                INDEX_STATUS['previewStage'] = 'termReport'
                report = preview_term_report(all_roots, grade, name, number, campus)
                INDEX_STATUS['previewStage'] = 'done'
                self.reply(200, {'schools': schools, 'hokushin': north, 'termReport': report})
            except Exception as exc:
                self.reply(422, {'error': str(exc)[:200]})
            return
        save_match = re.fullmatch(r'/save/([a-f0-9]{32})', self.path)
        if save_match:
            if not self.allowed():
                self.send_error(403)
                return
            with LOCK:
                session = SESSIONS.get(save_match[1])
                if not session or session[0] < time.time():
                    self.reply(404, {'error': '資料の表示期限が切れました。もう一度作成してください。'})
                    return
                try:
                    saved = session[2].get('savedPath')
                    if not saved:
                        saved = str(save_bundle_to_onedrive(Path(session[1].name) / 'staff-bundle.pdf', session[3]))
                        session[2]['savedPath'] = saved
                except Exception as exc:
                    self.reply(422, {'error': str(exc)[:200]})
                    return
                cloud_synced = bool(session[2].get('cloudSynced'))
            if not cloud_synced:
                try:
                    cloud_synced = sync_bundle_to_cloud(Path(saved))
                except Exception as exc:
                    self.reply(422, {'error': str(exc)[:200]})
                    return
                with LOCK:
                    session[2]['cloudSynced'] = cloud_synced
            self.reply(200, {'folder': str(Path(saved).parent), 'filename': Path(saved).name,
                             'cloudSynced': cloud_synced})
            return
        if self.path != '/generate' or not self.allowed() or self.headers.get('Content-Type', '').split(';')[0] != 'application/json':
            self.send_error(403)
            return
        try:
            length = int(self.headers.get('Content-Length', '0'))
            if length < 1 or length > 4096:
                raise ValueError('入力が大きすぎます')
            payload = json.loads(self.rfile.read(length))
            folder, manifest = make_bundle(payload)
            token = uuid.uuid4().hex
            with LOCK:
                expired = [key for key, value in SESSIONS.items() if value[0] < time.time()]
                for key in expired:
                    SESSIONS.pop(key)[1].cleanup()
                SESSIONS[token] = (time.time() + 900, folder, manifest, str(payload['number']))
            self.reply(200, {**manifest, 'combinedUrl': f'/pdf/{token}/staff-bundle.pdf',
                             'guideUrl': f'/pdf/{token}/guide.pdf' if manifest['hasGuide'] else None,
                             'saveUrl': f'/save/{token}'})
        except Exception as exc:
            self.reply(422, {'error': str(exc)[:200]})

    def log_message(self, format, *args):
        pass  # Do not log student names or paths.


if __name__ == '__main__':
    def index_forever():
        while True:
            try:
                INDEX_STATUS.pop('error', None)
                build_hokushin_index()
            except Exception as exc:
                INDEX_STATUS.update({'running': False, 'error': str(exc)[:200]})
            time.sleep(3600)
    threading.Thread(target=index_forever, daemon=True).start()
    def cleanup_sessions():
        while True:
            time.sleep(60)
            with LOCK:
                expired = [key for key, value in SESSIONS.items() if value[0] < time.time()]
                for key in expired:
                    SESSIONS.pop(key)[1].cleanup()
    threading.Thread(target=cleanup_sessions, daemon=True).start()
    ThreadingHTTPServer(('127.0.0.1', PORT), Handler).serve_forever()
