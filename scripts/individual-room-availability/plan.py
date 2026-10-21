"""Prepare, but do not publish, classroom availability using the journal reader."""
from datetime import date, datetime, timedelta
import hashlib
import json
from pathlib import Path
import sys
import openpyxl

root = Path(__file__).resolve().parents[2] / '.private' / 'kobetsu-20261021'
legacy_dir = Path.home() / 'OneDrive' / 'デスクトップ' / '【完成版】授業日誌システム'
sys.path.insert(0, str(legacy_dir))
import export_by_grade_subject as journal

notion_rows = json.loads((root / 'notion-rows.json').read_text(encoding='utf-8'))
rooms = {1: '①', 2: '②', 3: '③', 4: '④', 6: '⑥'}
session_columns = [('16:45', '18:15', range(6, 12)), ('18:35', '20:05', range(12, 19)), ('20:25', '21:55', range(19, 26))]

def jst(value):
    return datetime.fromisoformat(value.replace('Z', '+00:00'))

def overlap(row, start, end):
    value = row['date']
    if not value or not value.get('start'):
        return False
    raw = value['start']
    if 'T' not in raw:
        first = jst(raw + 'T00:00:00+09:00')
        last = jst((value.get('end') or raw) + 'T23:59:59+09:00')
    else:
        first = jst(raw)
        last = jst(value['end']) if value.get('end') else first + timedelta(minutes=90)
    return first < end and start < last

for month in ['2026-10', '2026-11']:
    source = root / 'legacy-reader' / month / f'{month[:4]}年{int(month[5:])}月スケジュール.xlsm'
    book = openpyxl.load_workbook(source, read_only=False, data_only=True, keep_vba=True)
    sheet = next(sh for campus, name, sh in journal.choose_target_sheets(book) if campus == '本校')
    classroom_row = journal.detect_classroom_row(sheet)
    all_candidates, blocked, existing, review = [], [], [], []
    for r in range(12, 50):
        if not isinstance(sheet.cell(r, 1).value, int):
            continue
        day, weekday = journal.parse_day_week_from_row(sheet, r)
        if not day or not str(day).isdigit():
            continue
        d = f'{month}-{int(day):02d}'
        actual_date = date.fromisoformat(d)
        if actual_date.weekday() == 6:
            continue
        if month == '2026-10' and actual_date < date(2026, 10, 21):
            continue
        if weekday not in ['月', '火', '水', '木', '金', '土']:
            raise RuntimeError(f'Unexpected weekday: {d}')
        notes = [journal.merged_value_text(sheet, r, c) for c in range(3, 26)]
        uncertain = any('後日' in x and ('対策' in x or '授業' in x) for x in notes)
        for start_clock, end_clock, columns in session_columns:
            start = jst(f'{d}T{start_clock}:00+09:00')
            end = jst(f'{d}T{end_clock}:00+09:00')
            for c in columns:
                room_text = journal.merged_value_text(sheet, classroom_row, c)
                number = next((n for n, symbol in rooms.items() if symbol in room_text), None)
                if number is None:
                    continue
                item = {'date': d, 'start': start_clock, 'end': end_clock, 'room': '本' + rooms[number],
                        'source_cell': sheet.cell(r, c).coordinate, 'title': '個別使用可能'}
                cell_value = journal.merged_value_text(sheet, r, c).strip()
                if cell_value:
                    blocked.append({**item, 'reason': '原表に記入あり', 'source_text': cell_value})
                    continue
                if uncertain:
                    review.append({**item, 'reason': 'オンライン期末対策の詳細が後日案内のため使用教室・時間帯未確定'})
                    continue
                matching = []
                uncertain_rows = []
                nonblocking_rows = []
                diagnostic_in_session = any('診断' in journal.merged_value_text(sheet, r, cc) for cc in columns)
                for row in notion_rows:
                    if not overlap(row, start, end):
                        continue
                    if not row['room'] and '本校' in row['campus']:
                        # Staff leave does not occupy a classroom.
                        # A diagnostic noted in the workbook already blocks its explicit classroom cell.
                        if ('休み' in row['tags'] and row['teachers']) or ('診断テスト' in row['tags'] and diagnostic_in_session):
                            nonblocking_rows.append({'id': row['id'], 'edited': row['edited']})
                            continue
                    if row['room'] == item['room']:
                        matching.append(row)
                    elif not row['room'] and '本校' in row['campus'] and any(tag in row['tags'] for tag in ['個別指導', '面談(対面)', '診断テスト', '入会手続き', '見学打合せ', '休み']):
                        symbols = [symbol for symbol in rooms.values() if symbol in row['title']]
                        if not symbols or rooms[number] in symbols:
                            uncertain_rows.append(row)
                if matching:
                    target = existing if all('個別使用可能' in row['tags'] for row in matching) else blocked
                    target.append({**item, 'reason': 'Notionに既存枠または使用予定あり', 'page_ids': [row['id'] for row in matching]})
                elif uncertain_rows:
                    review.append({**item, 'reason': '同時間帯に本校の教室未指定の予定あり', 'page_ids': [row['id'] for row in uncertain_rows]})
                else:
                    if nonblocking_rows:
                        item['resolved_unassigned_rows'] = nonblocking_rows
                    all_candidates.append(item)
    keys = [(x['date'], x['start'], x['room']) for x in all_candidates]
    assert len(keys) == len(set(keys)), 'Duplicate planned slot'
    assert all(x['room'] in ['本①', '本②', '本③', '本④', '本⑥'] for x in all_candidates)
    plan = {'month': month, 'applied': False, 'source': str(source), 'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
            'basis': '授業日誌のexport_by_grade_subject/collect_events・教室列・原表の記入・Notion既存予定',
            'target_month_confirmed': True, 'authorization': 'ユーザーが10月の残りと11月の両方を指定', 'candidates': all_candidates, 'existing': existing, 'blocked': blocked, 'review': review}
    (root / f'plan-{month}.json').write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'month': month, 'create_candidates': len(all_candidates), 'existing': len(existing), 'blocked': len(blocked), 'review': len(review),
                      'review_dates': sorted(set(x['date'] for x in review)), 'examples': all_candidates[:4]}, ensure_ascii=False, indent=2))
